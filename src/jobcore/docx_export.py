"""CV markdown -> .docx. Companion to pdf.py — same md_content contract (H1 name,
H2 sections in upper-case, H3 job entries, bullet lists, inline **bold**/*italic*/
[link](url)) and the same output-naming rules, reused directly from pdf.py so the
two renderers can never drift on what "CV_<Name>_<Company>" means.
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .console import force_utf8_output
from .pdf import _derive_stem, _strip_metadata

_DARK_BLUE = RGBColor(0x1A, 0x2E, 0x44)
_LINK_BLUE = "1A5276"

# Mirrors pdf.py's CV_CSS numbers exactly (A4, 12mm/14mm margins, 10pt body,
# 17/11/10pt headings, tight paragraph/list spacing) — same content, same target
# page count, just OOXML properties instead of CSS. Word's default styles (26pt
# Title, 14pt Heading 1, 24pt space-before) are what pushed this to 3 pages; the
# values below are the reason to override every style explicitly rather than rely
# on python-docx's defaults.
_BODY_PT = 10
_TAGLINE_PT = 9.5
_H1_PT = 17  # the candidate's name
_H2_PT = 11  # section headers (PROFESSIONAL SUMMARY, etc.)
_H3_PT = 10  # job entries (Company | Title | Dates)

# Order matters: **bold** must match before a lone *italic* does.
_INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]+\]\([^)]+\))")
_LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def _add_hyperlink(paragraph, url: str, text: str) -> None:
    """Insert a real, clickable hyperlink run — python-docx has no public API for
    this, so it's built directly via the OOXML relationship + <w:hyperlink> element.
    """
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), _LINK_BLUE)
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)

    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_inline_runs(paragraph, text: str) -> None:
    """Parse **bold**, *italic*, and [text](url) inline markdown into runs on an
    existing paragraph — the docx equivalent of what the `markdown` library does
    for the PDF/HTML path in pdf.py.
    """
    pos = 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos : m.start()])
        token = m.group(0)
        if token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("["):
            link_m = _LINK_RE.match(token)
            if link_m:
                _add_hyperlink(paragraph, link_m.group(2), link_m.group(1))
            else:  # unreachable (_INLINE_RE's link branch == _LINK_RE); narrows the Optional for mypy
                paragraph.add_run(token)
        else:  # single-asterisk italic
            paragraph.add_run(token[1:-1]).italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _set_style(style, *, size_pt, bold=None, before_pt=0, after_pt=0, line=1.05, keep_with_next=False):
    style.font.size = Pt(size_pt)
    if bold is not None:
        style.font.bold = bold
    pf = style.paragraph_format
    pf.space_before = Pt(before_pt)
    pf.space_after = Pt(after_pt)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line
    pf.keep_with_next = keep_with_next


def _setup_styles(doc) -> None:
    section = doc.sections[0]
    section.page_height = Mm(297)
    section.page_width = Mm(210)
    section.top_margin = Mm(12)
    section.bottom_margin = Mm(12)
    section.left_margin = Mm(14)
    section.right_margin = Mm(14)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    _set_style(normal, size_pt=_BODY_PT, after_pt=4, line=1.15)

    title = doc.styles["Title"]
    _set_style(title, size_pt=_H1_PT, bold=True, after_pt=1, line=1.0, keep_with_next=True)

    h1 = doc.styles["Heading 1"]
    _set_style(h1, size_pt=_H2_PT, bold=True, before_pt=8, after_pt=2, line=1.0, keep_with_next=True)

    h2 = doc.styles["Heading 2"]
    _set_style(h2, size_pt=_H3_PT, bold=True, before_pt=5, after_pt=1, line=1.0, keep_with_next=True)

    bullet = doc.styles["List Bullet"]
    bullet.font.name = "Calibri"
    _set_style(bullet, size_pt=_BODY_PT, after_pt=1, line=1.12)


def md_to_docx(md_content: str, out_path: Path) -> Path:
    doc = Document()
    _setup_styles(doc)

    after_h1 = False  # the tagline paragraph right under the name renders smaller/tighter
    for line in md_content.splitlines():
        stripped = line.strip()
        if not stripped or stripped == "---":
            continue

        if stripped.startswith("### "):
            p = doc.add_heading(level=2)
            _add_inline_runs(p, stripped[4:].strip())
            after_h1 = False
        elif stripped.startswith("## "):
            p = doc.add_heading(level=1)
            run = p.add_run(stripped[3:].strip().upper())
            run.font.color.rgb = _DARK_BLUE
            after_h1 = False
        elif stripped.startswith("# "):
            p = doc.add_heading(level=0)
            run = p.add_run(stripped[2:].strip())
            run.font.color.rgb = _DARK_BLUE
            after_h1 = True
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            _add_inline_runs(p, stripped[2:].strip())
            after_h1 = False
        else:
            p = doc.add_paragraph()
            _add_inline_runs(p, stripped)
            if after_h1:
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(3)
                for run in p.runs:
                    run.font.size = Pt(_TAGLINE_PT)
                after_h1 = False

    doc.save(str(out_path))
    return out_path


def render(md_path: Path, out_path: Path | None = None) -> Path:
    """Render a CV markdown file to .docx — same naming rule as pdf.render()."""
    md_path = Path(md_path)
    md_content = _strip_metadata(md_path.read_text(encoding="utf-8"))

    if out_path is None:
        out_path = md_path.parent / f"{_derive_stem(md_path, md_content)}.docx"
    out_path = Path(out_path)

    md_to_docx(md_content, out_path)
    print(f"DOCX saved: {out_path}")
    return out_path


if __name__ == "__main__":
    import sys

    force_utf8_output()
    if len(sys.argv) < 2:
        print("Usage: python -m src.jobcore.docx_export <CV.md> [out.docx]")
        sys.exit(1)
    render(Path(sys.argv[1]), Path(sys.argv[2]) if len(sys.argv) > 2 else None)
